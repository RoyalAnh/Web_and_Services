from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = REPORT_DIR.parent
REPOSITORY_DIR = PROJECT_DIR.parent
MARKDOWN_PATH = REPOSITORY_DIR / "Ex3-RBAC-2025229M-PhamQuangAnh.md"
TEMPLATE_PATH = REPOSITORY_DIR / "Example-ex3" / "Ex3-RBAC-MaSV-HoTenSV.docx"
OUTPUT_PATH = REPOSITORY_DIR / "Ex3-RBAC-2025229M-PhamQuangAnh.docx"


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def plain_text(value):
    value = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    value = value.replace("**", "").replace("__", "")
    value = value.replace("`", "").replace("\\$", "$")
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", value)
    return value.strip()


def parse_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line):
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def style_by_id(document, style_id, style_type):
    style = document.styles.get_by_id(style_id, style_type)
    if style is None:
        raise KeyError(f"Required Word style is missing: {style_id}")
    return style


def format_table(table, header=True):
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    if header and row_index == 0:
                        run.bold = True
            if header and row_index == 0:
                set_cell_shading(cell, "DCEAF2")


def add_markdown_table(document, lines):
    rows = [parse_table_row(line) for line in lines if not is_separator(line)]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            text = plain_text(row[column_index]) if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = True
    format_table(table, header=len(rows) > 1)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(55, 65, 75)
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:shd"))
    paragraph._p.pPr[-1].set(qn("w:fill"), "F2F5F7")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build():
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Word template not found: {TEMPLATE_PATH}")
    markdown = MARKDOWN_PATH.read_text(encoding="utf-8")
    document = Document(TEMPLATE_PATH)
    clear_body(document)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in [
        ("Title", 21, "143D59"),
        ("Heading 1", 16, "143D59"),
        ("Heading 2", 13, "1677A8"),
        ("Heading 3", 11, "1677A8"),
    ]:
        style = style_by_id(document, name.replace(" ", ""), WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped or stripped in {"---", "***"}:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = plain_text(heading.group(2))
            if level == 1 and text in {
                "ĐẠI HỌC BÁCH KHOA HÀ NỘI",
                "BÁO CÁO CÁ NHÂN",
                "CAM KẾT",
            }:
                paragraph = document.add_paragraph(
                    style=style_by_id(document, "Title", WD_STYLE_TYPE.PARAGRAPH)
                )
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
            elif level == 1:
                document.add_paragraph(
                    text,
                    style=style_by_id(document, "Heading1", WD_STYLE_TYPE.PARAGRAPH),
                )
            else:
                document.add_paragraph(
                    text,
                    style=style_by_id(
                        document, f"Heading{level}", WD_STYLE_TYPE.PARAGRAPH
                    ),
                )
            index += 1
            continue
        if stripped.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(
                style=style_by_id(document, "ListBullet", WD_STYLE_TYPE.PARAGRAPH)
            )
            paragraph.add_run(plain_text(stripped[2:]))
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.add_run(plain_text(stripped))
        index += 1

    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = "Báo cáo cá nhân — Fullstack RBAC"
    document.core_properties.subject = "Báo cáo theo template IT4409"
    document.save(OUTPUT_PATH)

    reopened = Document(OUTPUT_PATH)
    if len(reopened.paragraphs) < 50 or len(reopened.tables) < 20:
        raise RuntimeError("Generated DOCX is missing expected report structure.")
    print(f"Created readable DOCX: {OUTPUT_PATH} ({len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables)")


if __name__ == "__main__":
    build()
